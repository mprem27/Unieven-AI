"""Bounded, non-executing seating imports. Publishing requires staff review."""
import csv,io,re,zipfile
MAX_BYTES=8*1024*1024
FIELDS={'vtuNumber':{'vtu','vtunumber','vtuno','rollno','rollnumber','registerno','registernumber'},'hall':{'hall','block','building','examcentre'},'room':{'room','roomno','roomnumber','examroom'},'seat':{'seat','seatno','seatnumber','desk','deskno'},'reportingTime':{'reportingtime','reporttime','reporting'},'notes':{'notes','instructions','remarks'}}
VTU=re.compile(r'^VTU[0-9]{3,12}$',re.I)
def _heading(value):return re.sub(r'[^a-z0-9]','',str(value or '').lower())
def _cell(value):
    text=str(value if value is not None else '').strip()
    if len(text)>500 or '\x00' in text or text.startswith(('=','+','-','@')):raise ValueError('A seating cell is too long or contains a spreadsheet formula.')
    return text

def _table(lines):
    mapping=None;rows=[];warnings=[];seen=set()
    for index,line in enumerate(lines):
        if index>5100:raise ValueError('Use at most 5,000 student rows.')
        if not any(str(cell or '').strip() for cell in line):continue
        line=[_cell(cell) for cell in line]
        headers={}
        for position,value in enumerate(line):
            label=_heading(value)
            for field,aliases in FIELDS.items():
                if label in aliases:
                    if field in headers:raise ValueError('The table has ambiguous duplicate column labels.')
                    headers[field]=position
        if all(key in headers for key in ['vtuNumber','room','seat']):mapping=headers;continue
        if mapping is None:continue
        raw=line[mapping['vtuNumber']] if mapping['vtuNumber']<len(line) else ''
        number=str(raw or '').strip().upper()
        if not VTU.fullmatch(number):
            warnings.append(f'Row {index+1} has no exact VTU number; review it before publishing.');continue
        if number in seen:raise ValueError('Duplicate VTU number. Correct the source rows before publishing.')
        record={field:_cell(line[position]) if position<len(line) else '' for field,position in mapping.items()}
        record['vtuNumber']=number
        if not record.get('room') or not record.get('seat'):raise ValueError('Every VTU row needs a room and seat.')
        seen.add(number);rows.append({field:record.get(field,'') for field in FIELDS})
        if len(rows)>5000:raise ValueError('Use at most 5,000 student rows.')
    if mapping is None:warnings.append('No supported VTU / Room / Seat header was found. Enter or correct the rows in the review table.')
    if not rows:warnings.append('No student seats were extracted. For scanned PDFs, run OCR and review the output, or use the CSV template.')
    return rows,warnings[:20]

def _pdf(bytes_value):
    from pypdf import PdfReader
    from pypdf.generic import DictionaryObject,ArrayObject,IndirectObject
    if not bytes_value.startswith(b'%PDF-'):raise ValueError('The uploaded PDF does not match its format.')
    reader=PdfReader(io.BytesIO(bytes_value),strict=True)
    if reader.is_encrypted:raise ValueError('Encrypted PDFs are not supported.')
    if len(reader.pages)>60:raise ValueError('Use a PDF with at most 60 pages.')
    visited=set();objects=[reader.trailer];count=0
    while objects:
        item=objects.pop()
        if isinstance(item,IndirectObject):item=item.get_object()
        key=id(item)
        if key in visited:continue
        visited.add(key);count+=1
        if count>20000:raise ValueError('This PDF is too complex to import safely.')
        if isinstance(item,DictionaryObject):
            if any(str(key) in {'/JS','/JavaScript','/EmbeddedFiles','/Launch','/RichMedia'} for key in item) or str(item.get('/S')) in {'/JavaScript','/Launch'}:raise ValueError('PDF scripts or embedded files are not supported.')
            objects.extend(value for key,value in item.items() if str(key) not in {'/Parent'})
        elif isinstance(item,ArrayObject):objects.extend(item)
    text='\n'.join(page.extract_text() or '' for page in reader.pages)
    if len(text)>500000:raise ValueError('The PDF text is too large.')
    return [re.split(r'\s*\|\s*|\t+| {2,}',line.strip()) for line in text.splitlines() if line.strip()]

def _read_table_file(filename,content,allow_pdf=True):
    if not isinstance(content,bytes) or not content or len(content)>MAX_BYTES:raise ValueError('Upload one file between 1 byte and 8 MB.')
    extension=str(filename).lower().rsplit('.',1)[-1]
    try:
        if extension=='csv':
            text=content.decode('utf-8-sig')
            if '\x00' in text:raise ValueError('Use a UTF-8 CSV file.')
            try:delimiter=csv.Sniffer().sniff(text[:8192],delimiters=',;\t').delimiter
            except csv.Error:delimiter=','
            lines=[]
            for line in csv.reader(io.StringIO(text),delimiter=delimiter):
                if len(lines)>5100 or len(line)>40:raise ValueError('Use at most 5,000 rows and 40 columns.')
                lines.append(line)
            kind='csv'
        elif extension=='xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                files=archive.infolist()
                if len(files)>3000 or sum(item.file_size for item in files)>64*1024*1024 or any(item.file_size>24*1024*1024 for item in files):raise ValueError('The spreadsheet expands beyond the import limit.')
                if any('vbaProject' in item.filename or item.filename.startswith('xl/externalLinks/') for item in files):raise ValueError('Macros and external workbook links are not supported.')
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                if len(book.worksheets)!=1:raise ValueError('Use one seating worksheet per uploaded examination.')
                sheet=book.worksheets[0]
                if sheet.max_column>40 or sheet.max_row>5101:raise ValueError('Use at most 5,000 rows and 40 columns.')
                lines=[]
                for row in sheet.iter_rows():
                    values=[]
                    for cell in row:
                        if cell.data_type=='f':raise ValueError('Replace formulas with plain seating values.')
                        value=cell.value
                        if isinstance(value,int) and re.fullmatch(r'0{2,12}',cell.number_format or ''):value=str(value).zfill(len(cell.number_format))
                        values.append(value)
                    lines.append(values)
            finally:book.close()
            kind='xlsx'
        elif extension=='xls':raise ValueError('Legacy XLS is not supported by the installed parser. Save it as XLSX or CSV and upload again.')
        elif extension=='pdf' and allow_pdf:lines=_pdf(content);kind='pdf'
        else:raise ValueError('Upload a CSV, XLSX or text PDF seating list.')
        return lines,kind
    except ValueError:raise
    except Exception as exc:raise ValueError('This file could not be parsed. Use the supplied CSV template or enter rows manually.') from exc


def parse_seating_file(filename,content,course_code=None):
    lines,kind=_read_table_file(filename,content)
    if course_code is not None:
        return _course_table(lines,kind,course_code)
    rows,warnings=_table(lines)
    return {'rows':rows,'warnings':warnings,'format':kind,'requiresReview':True}


def _course_code(value):
    raw=_cell(value)
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._/ -]{0,39}',raw):raise ValueError('Use an exact course code for the roster filter.')
    return re.sub(r'[^A-Z0-9]','',raw.upper())


def parse_seating_roster(filename,content,course_code=None):
    """Read only a staff-selected roster; never infer registration or identities."""
    if str(filename).lower().rsplit('.',1)[-1] not in {'csv','xlsx'}:raise ValueError('Upload a CSV or XLSX student roster.')
    lines,kind=_read_table_file(filename,content,allow_pdf=False)
    selected=_course_code(course_code) if course_code is not None else None
    aliases={'vtuNumber':FIELDS['vtuNumber'],'name':{'name','studentname','fullname'},'courseCode':{'coursecode','subjectcode','papercode'}}
    mapping=None;students=[];seen=set();duplicates=0;data_rows=0
    for index,line in enumerate(lines):
        if not any(str(cell or '').strip() for cell in line):continue
        line=[_cell(cell) for cell in line]
        if mapping is None:
            headers={}
            for position,value in enumerate(line):
                for field,labels in aliases.items():
                    if _heading(value) in labels:
                        if field in headers:raise ValueError('The roster has ambiguous duplicate column labels.')
                        headers[field]=position
            if 'vtuNumber' not in headers:raise ValueError('The roster needs an exact VTU Number column. Download the roster template.')
            if selected and 'courseCode' not in headers:raise ValueError('A Course Code column is required when filtering the roster.')
            mapping=headers;continue
        data_rows+=1
        if data_rows>5000:raise ValueError('Use at most 5,000 student roster rows.')
        def value(field):
            position=mapping.get(field)
            return line[position] if position is not None and position<len(line) else ''
        if selected and _course_code(value('courseCode'))!=selected:continue
        number=value('vtuNumber').upper()
        if not VTU.fullmatch(number):raise ValueError(f'Roster row {index+1} needs an exact VTU number; no students were generated.')
        if number in seen:
            if selected:duplicates+=1;continue
            raise ValueError('Duplicate VTU number. Select an exact course filter or correct the roster.')
        seen.add(number);record={'vtuNumber':number}
        if value('name'):record['name']=value('name')
        students.append(record)
    if not students:raise ValueError('No roster students match the selected course code.' if selected else 'The roster has no students. Add an exact VTU Number column and student rows.')
    warnings=[]
    if duplicates:warnings.append(f'{duplicates} duplicate row(s) for the explicitly selected course were removed by exact VTU number.')
    return {'students':students,'warnings':warnings,'format':kind,'requiresReview':True}


def _course_table(lines,kind,course_code):
    selected=_course_code(course_code)
    aliases={**FIELDS,'name':{'name','studentname','fullname'},'courseCode':{'coursecode','subjectcode','papercode','subject'}}
    mapping=None;rows=[];issues=[];excluded=0
    for index,line in enumerate(lines):
        if index>5100:raise ValueError('Use at most 5,000 student rows.')
        if not any(str(cell or '').strip() for cell in line):continue
        line=[_cell(cell) for cell in line]
        headers={}
        for position,value in enumerate(line):
            for field,labels in aliases.items():
                if _heading(value) in labels:
                    if field in headers:raise ValueError('The table has ambiguous duplicate column labels.')
                    headers[field]=position
        if all(key in headers for key in ['room','seat']) and ('vtuNumber' in headers or 'name' in headers):
            mapping=headers;continue
        if mapping is None:continue
        def value(field):
            position=mapping.get(field)
            return line[position] if position is not None and position<len(line) else ''
        subject=value('courseCode')
        if not subject:
            issues.append({'sourceRow':index+1,'message':'Missing Course Code. Supply the exact examination subject.'});continue
        try:subject_key=_course_code(subject)
        except ValueError:
            issues.append({'sourceRow':index+1,'message':'Invalid Course Code.'});continue
        if subject_key!=selected:excluded+=1;continue
        number=value('vtuNumber').upper()
        if number and not VTU.fullmatch(number):
            issues.append({'sourceRow':index+1,'message':'Invalid supplied VTU; no name fallback attempted.'});continue
        if not number and not value('name'):
            issues.append({'sourceRow':index+1,'message':'An exact VTU or student name is required.'});continue
        if not value('room') or not value('seat'):
            issues.append({'sourceRow':index+1,'message':'Room and seat are required.'});continue
        rows.append({**{field:value(field) for field in FIELDS},'vtuNumber':number,'name':value('name'),'courseCode':subject,'sourceRow':index+1})
        if len(rows)>5000:raise ValueError('Use at most 5,000 student rows.')
    warnings=[]
    if excluded:warnings.append(f'{excluded} row(s) for other course codes were excluded by exact subject matching.')
    if mapping is None:warnings.append('No supported Course Code / VTU or Name / Room / Seat table was found.')
    if not rows:warnings.append('No matching seats were extracted. For scanned PDFs, run OCR and check the output, or upload the CSV template.')
    return {'rows':rows,'issues':issues,'warnings':warnings,'format':kind,'requiresReview':True}
