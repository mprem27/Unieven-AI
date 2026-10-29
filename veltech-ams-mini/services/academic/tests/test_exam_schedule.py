import base64
import importlib
import io
import unittest
import zipfile
from datetime import datetime, time
from pathlib import Path


def text_pdf(pages, script=None):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for items in pages:
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({
            NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica'),
        })
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): font}),
        })
        commands = []
        for x, y, value in items:
            escaped = value.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            commands.append(f'BT /F1 10 Tf 1 0 0 1 {x} {y} Tm ({escaped}) Tj ET')
        stream = DecodedStreamObject()
        stream.set_data('\n'.join(commands).encode('latin1'))
        page[NameObject('/Contents')] = stream
    if script:
        writer.add_js(script)
    data = io.BytesIO()
    writer.write(data)
    return data.getvalue()


class ExamScheduleParserTests(unittest.TestCase):
    def parser(self):
        try:
            module = importlib.import_module('app.services.exam_schedule')
        except ModuleNotFoundError:
            self.fail('The bounded examination schedule parser is missing.')
        return module.parse_exam_schedule_file

    def test_csv_extracts_schedule_fields_without_inventing_metadata(self):
        result = self.parser()('schedule.csv', (
            'Course Code,Course Title,Date,Start Time,End Time,Department,Semester,Section,Room\n'
            'CS-101,Discrete Mathematics,26.10.2026,4.30 pm,6.00 pm,CSE,7,A,B201\n'
        ).encode())
        self.assertEqual(result['rows'], [{
            'courseCode': 'CS-101', 'courseName': 'Discrete Mathematics',
            'date': '2026-10-26', 'startTime': '16:30', 'endTime': '18:00',
            'department': 'CSE', 'semester': '7', 'section': 'A', 'room': 'B201',
        }])
        self.assertIs(result['requiresReview'], True)
        self.assertIsInstance(result['warnings'], list)
        self.assertNotIn('title', result)

    def test_csv_quoted_titles_and_combined_time_ranges(self):
        result = self.parser()('exam.csv', (
            'CourseCode;CourseName;Exam Date;Time\n'
            'CS101;"Data, Structures";2026-11-02;09:00 - 10:30\n'
        ).encode())
        self.assertEqual(result['rows'][0]['courseName'], 'Data, Structures')
        self.assertEqual(result['rows'][0]['date'], '2026-11-02')
        self.assertEqual(result['rows'][0]['endTime'], '10:30')

    def test_csv_does_not_forward_fill_missing_dates_or_repair_ambiguous_codes(self):
        result = self.parser()('exam.csv', (
            'Course Code,Course Title,Date,Start Time,End Time\n'
            'CS101,Algorithms,2026-10-26,09:00,10:30\n'
            '10211CSI03,Operating Systems,,09:00,10:30\n'
            'CS103,Networks,31.02.2026,10:00,09:00\n'
        ).encode())
        self.assertEqual(result['rows'][1]['courseCode'], '10211CSI03')
        self.assertEqual(result['rows'][1]['date'], '')
        self.assertEqual(result['rows'][2]['date'], '')
        self.assertTrue(any('date' in warning.lower() for warning in result['warnings']))
        self.assertTrue(any('code' in warning.lower() for warning in result['warnings']))
        self.assertTrue(any('end' in warning.lower() for warning in result['warnings']))

    def test_csv_rejects_formulas_duplicate_headers_and_row_overflow(self):
        parse = self.parser()
        for content in [
            b'Course Code,Course Title,Date,Start Time,End Time\nCS1,=WEBSERVICE("bad"),2026-10-26,09:00,10:30',
            b'Course Code,CourseCode,Course Title,Date,Start Time,End Time\nCS1,CS2,Title,2026-10-26,09:00,10:30',
            b'Course Code,Course Title,Date,Start Time,End Time\n' + b'CS1,Title,2026-10-26,09:00,10:30\n' * 501,
        ]:
            with self.subTest(content=content[:70]):
                with self.assertRaises(ValueError):
                    parse('exam.csv', content)

    def test_xlsx_preserves_dates_times_and_explicit_merged_date_cells(self):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(['Course Code', 'Course Title', 'Date', 'Start Time', 'End Time'])
        sheet.append(['CS101', 'Algorithms', datetime(2026, 10, 26), time(16, 30), time(18)])
        sheet.append(['CS102', 'Networks', None, time(16, 30), time(18)])
        sheet.merge_cells('C2:C3')
        data = io.BytesIO()
        book.save(data)
        result = self.parser()('exam.xlsx', data.getvalue())
        self.assertEqual([row['date'] for row in result['rows']], ['2026-10-26'] * 2)
        self.assertEqual(result['rows'][1]['startTime'], '16:30')
        self.assertTrue(any('merged' in warning.lower() for warning in result['warnings']))

    def test_xlsx_rejects_formulas_multiple_sheets_and_external_links(self):
        from openpyxl import Workbook

        parse = self.parser()
        book = Workbook()
        book.active.append(['Course Code', 'Course Title', 'Date', 'Start Time', 'End Time'])
        book.active.append(['CS1', '=1+1', '2026-10-26', '09:00', '10:30'])
        data = io.BytesIO()
        book.save(data)
        with self.assertRaises(ValueError):
            parse('exam.xlsx', data.getvalue())
        book.active['B2'] = 'Title'
        book.create_sheet('Another')
        data = io.BytesIO()
        book.save(data)
        with self.assertRaises(ValueError):
            parse('exam.xlsx', data.getvalue())
        data = io.BytesIO()
        with zipfile.ZipFile(data, 'w') as archive:
            archive.writestr('xl/externalLinks/externalLink1.xml', '<externalLink/>')
        with self.assertRaises(ValueError):
            parse('exam.xlsx', data.getvalue())

    def test_spreadsheet_entities_and_expanded_archives_fail_closed(self):
        from openpyxl import Workbook

        parse = self.parser()
        book = Workbook()
        book.active.append(['Course Code', 'Course Title', 'Date', 'Start Time', 'End Time'])
        book.active.append(['CS101', 'Algorithms', '2026-10-26', '09:00', '10:30'])
        original = io.BytesIO()
        book.save(original)
        data = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(data, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for item in source.infolist():
                content = source.read(item)
                if item.filename == 'xl/workbook.xml':
                    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE workbook [<!ENTITY name "value">]>' + content.decode().replace('name="Sheet"', 'name="&name;"')
                    content = xml.encode('utf-16')
                archive.writestr(item.filename, content)
        with self.assertRaisesRegex(ValueError, 'entities'):
            parse('exam.xlsx', data.getvalue())
        data = io.BytesIO()
        with zipfile.ZipFile(data, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/worksheets/sheet1.xml', b' ' * (24 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(ValueError, 'expands'):
            parse('exam.xlsx', data.getvalue())

    def test_pdf_page_and_decompressed_stream_limits_fail_closed(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, NameObject

        parse = self.parser()
        writer = PdfWriter()
        for _ in range(31):
            writer.add_blank_page(600, 800)
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(ValueError, '30 pages'):
            parse('exam.pdf', data.getvalue())
        writer = PdfWriter()
        page = writer.add_blank_page(600, 800)
        stream = DecodedStreamObject()
        stream.set_data(b' ' * (8 * 1024 * 1024 + 1))
        page[NameObject('/Contents')] = stream.flate_encode()
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaises(ValueError):
            parse('exam.pdf', data.getvalue())

    def test_positioned_pdf_groups_centered_merged_dates_and_wraps_titles(self):
        content = text_pdf([[
            (60, 760, 'Test-II and Mid Term Test-I Schedule'),
            (390, 735, 'AN: 4.30 pm to 6.00 pm'),
            (60, 700, 'Course Code'), (180, 700, 'Course Title'),
            (400, 700, 'Date/Day of'), (520, 700, 'Session'),
            (60, 660, 'CS101'), (180, 660, 'Web and Mobile Application'),
            (180, 648, 'Development'),
            (60, 620, 'CS102'), (180, 620, 'Algorithms'),
            (400, 646, '26.10.2026/'), (400, 634, 'Monday'), (520, 640, 'AN'),
            (60, 580, 'CS103'), (180, 580, 'Networks'),
            (400, 586, '27.10.2026/'), (400, 574, 'Tuesday'), (520, 580, 'AN'),
        ]])
        result = self.parser()('Test-III-Schedule.pdf', content)
        self.assertEqual([row['date'] for row in result['rows']], ['2026-10-26', '2026-10-26', '2026-10-27'])
        self.assertEqual(result['rows'][0]['courseName'], 'Web and Mobile Application Development')
        self.assertEqual(result['rows'][2]['startTime'], '16:30')
        self.assertEqual(result['title'], 'Test-II and Mid Term Test-I Schedule')
        self.assertTrue(any('heading' in warning.lower() for warning in result['warnings']))

    def test_pdf_with_unknown_session_never_invents_times(self):
        content = text_pdf([[
            (60, 700, 'Course Code'), (180, 700, 'Course Title'),
            (400, 700, 'Date'), (520, 700, 'Session'),
            (60, 660, 'CS101'), (180, 660, 'Algorithms'),
            (400, 660, '26.10.2026'), (520, 660, 'AN'),
        ]])
        result = self.parser()('exam.pdf', content)
        self.assertEqual(result['rows'][0]['startTime'], '')
        self.assertEqual(result['rows'][0]['endTime'], '')
        self.assertTrue(any('time' in warning.lower() for warning in result['warnings']))

    def test_scanned_pdf_returns_empty_reviewable_result(self):
        result = self.parser()('scanned.pdf', text_pdf([[]]))
        self.assertEqual(result['rows'], [])
        self.assertTrue(result['warnings'])
        self.assertIs(result['requiresReview'], True)

    def test_unsafe_or_unbounded_documents_are_rejected(self):
        from pypdf import PdfWriter

        parse = self.parser()
        for filename, content in [
            ('exam.csv', b'x' * (8 * 1024 * 1024 + 1)),
            ('exam.xlsm', b'not supported'), ('exam.xlsx', b'not a zip'),
            ('exam.pdf', b'not PDF'),
            ('exam.pdf', text_pdf([[]], script='app.alert("blocked")')),
        ]:
            with self.subTest(filename=filename):
                with self.assertRaises(ValueError):
                    parse(filename, content)
        writer = PdfWriter()
        writer.add_blank_page(600, 800)
        writer.encrypt('private')
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaises(ValueError):
            parse('encrypted.pdf', data.getvalue())

    def test_pdf_private_capture_metadata_is_allowed_but_arbitrary_attachments_are_rejected(self):
        from pypdf import PdfWriter
        from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject, TextStringObject

        parse = self.parser()
        for payload, allowed in [(b'{"type":"Document","isBackSide":false}', True), (b'{"type":"Document","script":"bad"}', False)]:
            writer = PdfWriter()
            page = writer.add_blank_page(600, 800)
            stream = DecodedStreamObject()
            stream.set_data(payload)
            stream[NameObject('/Type')] = NameObject('/EmbeddedFile')
            stream[NameObject('/Subtype')] = NameObject('/application/json')
            specification = DictionaryObject({
                NameObject('/Type'): NameObject('/Filespec'),
                NameObject('/AFRelationship'): NameObject('/ADBE_Contact_Private'),
                NameObject('/F'): TextStringObject('pageEntities.json'),
                NameObject('/EF'): DictionaryObject({NameObject('/F'): writer._add_object(stream)}),
            })
            page[NameObject('/AF')] = ArrayObject([writer._add_object(specification)])
            data = io.BytesIO()
            writer.write(data)
            if allowed:
                self.assertEqual(parse('scan.pdf', data.getvalue())['rows'], [])
            else:
                with self.assertRaises(ValueError):
                    parse('scan.pdf', data.getvalue())
        writer = PdfWriter()
        writer.add_blank_page(600, 800)
        writer.add_attachment('payload.exe', b'dangerous attachment')
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaises(ValueError):
            parse('attached.pdf', data.getvalue())

    def test_supplied_pdf_keeps_all_35_rows_and_merged_dates_without_code_repairs(self):
        reference = Path('/tmp/codex-remote-attachments/01a13885-eee9-764c-95b6-93a5c15ed078/0d1585c1-54ef-4f4c-8837-9bc571b86de0/1-CSE-Final-Year-Test-III-and-Mid-Term-Test-II-Schedule-SS2627.pdf')
        if not reference.exists():
            self.skipTest('The user attachment is available only in the delivery workspace.')
        result = self.parser()(reference.name, reference.read_bytes())
        expected = ['2026-10-26'] * 5 + ['2026-10-27'] * 9 + ['2026-10-28', '2026-10-29'] + ['2026-10-30'] * 9 + ['2026-11-02'] * 10
        self.assertEqual(len(result['rows']), 35)
        self.assertEqual([row['date'] for row in result['rows']], expected)
        self.assertTrue(all(row['startTime'] == '16:30' and row['endTime'] == '18:00' for row in result['rows']))
        self.assertEqual(result['rows'][25]['courseCode'], '10211CSI03')
        self.assertEqual(result['rows'][30]['courseCode'], 'J0211CS208')
        self.assertTrue(any('OCR' in warning for warning in result['warnings']))
        self.assertTrue(any('heading' in warning.lower() for warning in result['warnings']))


class ExamScheduleRouteTests(unittest.TestCase):
    def setUp(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from app.config.settings import settings

        try:
            from app.routes.exam_schedule import router
        except ModuleNotFoundError:
            self.fail('The private examination parser bridge route is missing.')
        self.settings = settings
        self.previous_key = settings.campus_service_key
        settings.campus_service_key = 'e' * 64
        app = FastAPI()
        app.include_router(router)
        self.client = TestClient(app)
        self.headers = {'X-Campus-Service-Key': settings.campus_service_key}
        self.body = {
            'filename': 'exam.csv',
            'contentBase64': base64.b64encode(b'Course Code,Course Title,Date,Start Time,End Time\nCS1,Title,2026-10-26,09:00,10:30').decode(),
        }

    def tearDown(self):
        if hasattr(self, 'settings'):
            self.settings.campus_service_key = self.previous_key

    def test_private_service_key_is_required(self):
        for headers in [{}, {'X-Campus-Service-Key': 'wrong'}]:
            response = self.client.post('/campus/bridge/exams/parse', json=self.body, headers=headers)
            self.assertEqual(response.status_code, 403)

    def test_route_returns_reviewable_rows_and_rejects_malformed_uploads(self):
        response = self.client.post('/campus/bridge/exams/parse', json=self.body, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['rows'][0]['courseCode'], 'CS1')
        self.assertIs(response.json()['requiresReview'], True)
        for body in [{**self.body, 'contentBase64': 'not!base64'}, {**self.body, 'filename': 'exam.exe'}]:
            self.assertEqual(self.client.post('/campus/bridge/exams/parse', json=body, headers=self.headers).status_code, 400)
        self.assertEqual(self.client.post('/campus/bridge/exams/parse', json={**self.body, 'published': True}, headers=self.headers).status_code, 422)

    def test_route_is_registered_under_the_application_api_prefix(self):
        from app.main import app
        from fastapi.testclient import TestClient

        response = TestClient(app).post('/api/v1/campus/bridge/exams/parse', json=self.body, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertIs(response.json()['requiresReview'], True)
